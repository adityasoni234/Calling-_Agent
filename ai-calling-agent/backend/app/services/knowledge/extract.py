"""Turn an uploaded file into plain text."""

from __future__ import annotations

import io

from app.core.errors import ValidationFailedError

SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text",
    "text/markdown": "text",
    "text/x-markdown": "text",
}
EXTENSION_MAP = {".pdf": "pdf", ".docx": "docx", ".txt": "text", ".md": "text", ".markdown": "text"}


def detect_kind(filename: str, mime_type: str | None) -> str:
    if mime_type and mime_type in SUPPORTED:
        return SUPPORTED[mime_type]
    lowered = filename.lower()
    for extension, kind in EXTENSION_MAP.items():
        if lowered.endswith(extension):
            return kind
    raise ValidationFailedError(
        "Unsupported file type. Upload PDF, DOCX, TXT, or Markdown.",
        code="unsupported_file_type",
    )


def extract_text(data: bytes, *, filename: str, mime_type: str | None = None) -> str:
    kind = detect_kind(filename, mime_type)

    if kind == "text":
        return data.decode("utf-8", errors="replace")

    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)
